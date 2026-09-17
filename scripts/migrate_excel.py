import os
import json
import re
from datetime import datetime
import openpyxl

def clean_str(val):
    if val is None:
        return ''
    s = str(val).strip()
    return re.sub(r'\s+', ' ', s)

def clean_num(val):
    if val is None or val == '':
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(',', '')
    if '+' in s:
        try:
            return float(sum(float(x.strip()) for x in s.split('+') if x.strip()))
        except:
            pass
    try:
        return float(s)
    except:
        return 0.0

def parse_excel():
    print("Loading 'عملاء اجل الاقساط .xlsx'...")
    wb_cust = openpyxl.load_workbook('عملاء اجل الاقساط .xlsx', data_only=True)
    
    customers_map = {}
    contracts = []
    suppliers = []

    # 1. Parse Letter Sheets for Customer Cards
    for sname in wb_cust.sheetnames:
        if 'مورد' in sname or 'تقرير' in sname:
            continue
        sheet = wb_cust[sname]
        max_r = sheet.max_row or 200
        
        r = 1
        while r <= max_r:
            c1 = clean_str(sheet.cell(r, 1).value)
            c2 = clean_str(sheet.cell(r, 2).value)
            c3 = clean_str(sheet.cell(r, 3).value)
            
            cust_name = ''
            card_start = None
            
            if 'اسم العميل' in c1:
                cust_name = c2 or c3
                card_start = r
            elif 'اسم العميل' in c2:
                cust_name = c3 or clean_str(sheet.cell(r, 4).value)
                card_start = r
            
            if card_start and cust_name and cust_name != '0' and len(cust_name) > 1 and not any(k in cust_name for k in ['سعر', 'البيان', 'مقدم', 'قسط']):
                device_name = ''
                phone = ''
                cash_price = 0.0
                inst_price = 0.0
                down_payment = 0.0
                down_date = ''
                remaining_bal = 0.0
                
                # Check next 3 rows for device, phone, prices
                for offset in range(1, 4):
                    row_idx = card_start + offset
                    if row_idx > max_r:
                        break
                    line_vals = [clean_str(sheet.cell(row_idx, c).value) for c in range(1, 10)]
                    raw_vals = [sheet.cell(row_idx, c).value for c in range(1, 10)]
                    
                    for idx, val in enumerate(line_vals):
                        if 'سعر الكاش' in val and idx + 1 < len(raw_vals):
                            cash_price = clean_num(raw_vals[idx + 1])
                        if 'سعر القسط' in val and idx + 1 < len(raw_vals):
                            inst_price = clean_num(raw_vals[idx + 1])
                        if re.match(r'^0?1[0125]\d{8}$', val) or ('10' in val and len(val) >= 9 and val.isdigit()):
                            phone = val
                    
                    cand_device = line_vals[1] if len(line_vals) > 1 else ''
                    if not cand_device and len(line_vals) > 2:
                        cand_device = line_vals[2]
                    if cand_device and not any(k in cand_device for k in ['رقم', 'سعر', 'نوع', 'البيان', 'مقدم', 'قسط']):
                        if not device_name:
                            device_name = cand_device

                installments_list = []
                inst_r = card_start + 4
                while inst_r <= min(card_start + 22, max_r):
                    cell_b = sheet.cell(inst_r, 2).value
                    cell_c = clean_str(sheet.cell(inst_r, 3).value)
                    cell_d = clean_str(sheet.cell(inst_r, 4).value)
                    cell_e = clean_str(sheet.cell(inst_r, 5).value)
                    cell_f = sheet.cell(inst_r, 6).value
                    
                    if 'المجموع' in cell_c or 'اجمالي' in cell_c or 'المتبقي' in clean_str(sheet.cell(inst_r, 5).value):
                        if clean_str(sheet.cell(inst_r, 5).value) == 'المتبقي' or 'المتبقي' in clean_str(sheet.cell(inst_r, 5).value):
                            remaining_bal = clean_num(cell_f)
                        inst_r += 1
                        continue
                    
                    if 'اسم العميل' in clean_str(cell_b) or 'اسم العميل' in cell_c:
                        break
                        
                    amt = clean_num(cell_b)
                    if amt > 0 or cell_c in ['مقدم', 'قسط'] or cell_d:
                        due_amt = amt
                        paid_amt = clean_num(cell_f) if cell_f is not None else 0.0
                        due_date = cell_d
                        paid_date = cell_e
                        item_type = 'down_payment' if 'مقدم' in cell_c else 'installment'
                        
                        if item_type == 'down_payment':
                            down_payment = due_amt
                            down_date = due_date or paid_date
                        
                        status = 'paid' if paid_amt >= due_amt and due_amt > 0 else ('partially_paid' if paid_amt > 0 else 'pending')
                        
                        installments_list.append({
                            'item_type': item_type,
                            'due_amount': due_amt,
                            'due_date': due_date,
                            'paid_date': paid_date,
                            'paid_amount': paid_amt,
                            'remaining_amount': max(due_amt - paid_amt, 0.0),
                            'status': status
                        })
                    inst_r += 1

                if remaining_bal == 0.0 and inst_price > 0:
                    tot_paid = sum(i['paid_amount'] for i in installments_list)
                    remaining_bal = max(inst_price - tot_paid, 0.0)

                if cust_name not in customers_map:
                    customers_map[cust_name] = {
                        'name': cust_name,
                        'phone': phone,
                        'sheet': sname,
                        'total_contracts': 0,
                        'contracts': []
                    }
                
                contract_data = {
                    'customer_name': cust_name,
                    'device_name': device_name or 'جهاز هاتف ذكي',
                    'cash_price': cash_price,
                    'installment_price': inst_price,
                    'down_payment': down_payment,
                    'down_payment_date': down_date,
                    'remaining_balance': remaining_bal,
                    'installment_count': len([i for i in installments_list if i['item_type'] == 'installment']),
                    'installments': installments_list
                }
                customers_map[cust_name]['total_contracts'] += 1
                customers_map[cust_name]['contracts'].append(contract_data)
                contracts.append(contract_data)
                
                r = inst_r
                continue
            r += 1

    # 2. Parse Suppliers Sheet (Headers are in Row 2)
    for sname in wb_cust.sheetnames:
        if 'مورد' in sname:
            sup_sheet = wb_cust[sname]
            for col_idx in range(1, (sup_sheet.max_column or 30) + 1):
                s_name = clean_str(sup_sheet.cell(2, col_idx).value)
                if s_name and s_name not in [s['name'] for s in suppliers] and len(s_name) > 2:
                    suppliers.append({
                        'name': s_name,
                        'notes': 'مورد مسجل من شيت الموردين'
                    })

    wb_cust.close()

    # 3. Parse Fast Credit from Journal (Headers are in Row 2)
    fast_credit_accounts = []
    print("Loading 'دفتر يومية  2025.xlsx'...")
    wb_jour = openpyxl.load_workbook('دفتر يومية  2025.xlsx', data_only=True)
    for sname in wb_jour.sheetnames:
        if 'اجل سريع' in sname:
            fc_sheet = wb_jour[sname]
            for c in range(2, (fc_sheet.max_column or 30) + 1):
                acc_name = clean_str(fc_sheet.cell(2, c).value)
                if acc_name and not any(k in acc_name for k in ['البيان', 'المبلغ', 'التاريخ', 'None']) and len(acc_name) > 1:
                    if acc_name not in [a['name'] for a in fast_credit_accounts]:
                        fast_credit_accounts.append({
                            'name': acc_name,
                            'type': 'partner_shop'
                        })
    wb_jour.close()

    # Summary report
    tot_contract_val = sum(c['installment_price'] for c in contracts)
    tot_rem_val = sum(c['remaining_balance'] for c in contracts)
    
    print(f"\n==================================================")
    print(f"MIGRATION EXTRACTION SUMMARY")
    print(f"==================================================")
    print(f"Total Unique Customers Found: {len(customers_map)}")
    print(f"Total Installment Contracts Found: {len(contracts)}")
    print(f"Total Contracts Value: {tot_contract_val:,.2f} EGP")
    print(f"Total Remaining Balance: {tot_rem_val:,.2f} EGP")
    print(f"Suppliers Found: {len(suppliers)} -> {[s['name'] for s in suppliers]}")
    print(f"Fast Credit Partner Accounts: {len(fast_credit_accounts)} -> {[a['name'] for a in fast_credit_accounts]}")

    output_data = {
        'extracted_at': datetime.now().isoformat(),
        'summary': {
            'total_customers': len(customers_map),
            'total_contracts': len(contracts),
            'total_contract_value': tot_contract_val,
            'total_remaining_balance': tot_rem_val,
            'suppliers_count': len(suppliers),
            'fast_credit_count': len(fast_credit_accounts)
        },
        'customers': list(customers_map.values()),
        'suppliers': suppliers,
        'fast_credit_accounts': fast_credit_accounts
    }

    with open('migrated_data.json', 'w', encoding='utf-8') as f:
        json.dump(output_data, f, ensure_ascii=False, indent=2)
    print("Migration data saved to 'migrated_data.json'.")

    # Generate SQL seed file for Supabase
    generate_sql_seed(output_data)

def generate_sql_seed(data):
    org_id = '00000000-0000-0000-0000-000000000001'
    treasury_id = '00000000-0000-0000-0000-000000000002'
    lines = []
    lines.append("-- ==============================================================")
    lines.append("-- Seed Data Migrated from Excel into Supabase PostgreSQL")
    lines.append("-- ==============================================================\n")

    lines.append(f"""
-- 1. المنشأة والخزينة الرئيسية
INSERT INTO public.organizations (id, name, trade_name, currency, currency_symbol)
VALUES ('{org_id}', 'سنترال المركزي', 'المركزي للمبيعات والأقساط', 'EGP', 'ج.م')
ON CONFLICT (id) DO UPDATE SET name = EXCLUDED.name;

INSERT INTO public.treasuries (id, organization_id, name, treasury_type, opening_balance, current_balance)
VALUES ('{treasury_id}', '{org_id}', 'الدرج الرئيسي', 'drawer', 0, 0)
ON CONFLICT (id) DO NOTHING;

-- 2. ماكينات الدفع الإلكتروني (فوري، أمان، بساطة، أمان تاتش)
INSERT INTO public.pos_machines (organization_id, name) VALUES
('{org_id}', 'فوري'),
('{org_id}', 'أمان'),
('{org_id}', 'بساطة'),
('{org_id}', 'أمان تاتش')
ON CONFLICT DO NOTHING;

-- 3. خطوط ومحافظ الكاش الموثقة من الدفتر
INSERT INTO public.cash_wallets (organization_id, phone_number, provider, account_label) VALUES
('{org_id}', '01091576032', 'vodafone_cash', 'خط كاش 1'),
('{org_id}', '01002919441', 'vodafone_cash', 'خط كاش 2'),
('{org_id}', '01033307821', 'vodafone_cash', 'خط كاش 3'),
('{org_id}', '01098968373', 'vodafone_cash', 'خط كاش 4'),
('{org_id}', '01090067941', 'vodafone_cash', 'خط كاش 5'),
('{org_id}', '01005168098', 'vodafone_cash', 'خط كاش 6')
ON CONFLICT DO NOTHING;
""")

    # 4. الموردين
    lines.append("\n-- الموردين المرحلين من الإكسل")
    for sup in data['suppliers']:
        sname_esc = sup['name'].replace("'", "''")
        lines.append(f"INSERT INTO public.suppliers (organization_id, name, notes) VALUES ('{org_id}', '{sname_esc}', 'مورد مرحّل من الإكسل') ON CONFLICT DO NOTHING;")

    # 5. حسابات الأجل السريع
    lines.append("\n-- حسابات الأجل السريع (المحلات الزميلة)")
    for fc in data['fast_credit_accounts']:
        fname_esc = fc['name'].replace("'", "''")
        lines.append(f"INSERT INTO public.fast_credit_accounts (organization_id, name, account_type) VALUES ('{org_id}', '{fname_esc}', 'partner_shop') ON CONFLICT DO NOTHING;")

    # 6. العملاء وعقود الأقساط
    lines.append("\n-- العملاء وعقود الأقساط وجداول السداد")
    cust_idx = 1
    for c in data['customers']:
        cname_esc = c['name'].replace("'", "''")
        phone_esc = (c['phone'] or '').replace("'", "''")
        code = f"CUS-{cust_idx:05d}"
        
        lines.append(f"""
DO $$
DECLARE
    v_cust_id UUID;
    v_contract_id UUID;
BEGIN
    INSERT INTO public.customers (organization_id, code, name, phone)
    VALUES ('{org_id}', '{code}', '{cname_esc}', NULLIF('{phone_esc}', ''))
    RETURNING id INTO v_cust_id;
""")
        ctr_idx = 1
        for ctr in c['contracts']:
            ctr_code = f"CTR-{cust_idx:04d}-{ctr_idx}"
            dev_esc = ctr['device_name'].replace("'", "''")
            cash_p = ctr['cash_price']
            inst_p = ctr['installment_price']
            down_p = ctr['down_payment']
            rem_b = ctr['remaining_balance']
            cnt = ctr['installment_count']
            
            lines.append(f"""
    INSERT INTO public.contracts (
        organization_id, customer_id, contract_number, device_name,
        cash_price, total_installment_price, down_payment, remaining_balance,
        installment_count, monthly_installment_amount, start_date, status
    ) VALUES (
        '{org_id}', v_cust_id, '{ctr_code}', '{dev_esc}',
        {cash_p}, {inst_p}, {down_p}, {rem_b},
        {cnt}, {rem_b / max(cnt, 1):.2f}, CURRENT_DATE,
        CASE WHEN {rem_b} <= 0 THEN 'completed' ELSE 'active' END
    ) RETURNING id INTO v_contract_id;
""")
            for inst_no, inst in enumerate(ctr['installments'], 1):
                d_amt = inst['due_amount']
                p_amt = inst['paid_amount']
                r_amt = inst['remaining_amount']
                st = inst['status']
                lines.append(f"""
    INSERT INTO public.installments (
        organization_id, contract_id, customer_id, installment_number,
        due_date, due_amount, paid_amount, remaining_amount, status
    ) VALUES (
        '{org_id}', v_contract_id, v_cust_id, {inst_no},
        CURRENT_DATE, {d_amt}, {p_amt}, {r_amt}, '{st}'
    );
""")
            ctr_idx += 1
            
        lines.append("END $$;\n")
        cust_idx += 1

    with open('seed_data.sql', 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print("Complete SQL Seed file generated: 'seed_data.sql'.")

if __name__ == '__main__':
    parse_excel()
